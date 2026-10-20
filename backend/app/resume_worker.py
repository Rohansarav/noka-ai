"""Isolated text-only parser. No files are extracted, executed or retained."""
import io
import json
import sys
import zipfile

MAX_TEXT = 20000


def bound_memory():
    """Limit the isolated parser to 512 MiB; fail closed if the OS cannot apply it."""
    if sys.platform != "win32":
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
        return None
    import ctypes
    from ctypes import wintypes
    class BasicLimits(ctypes.Structure):
        _fields_ = [("process_time", ctypes.c_longlong), ("job_time", ctypes.c_longlong), ("flags", wintypes.DWORD),
                    ("min_working_set", ctypes.c_size_t), ("max_working_set", ctypes.c_size_t), ("active_processes", wintypes.DWORD),
                    ("affinity", ctypes.c_size_t), ("priority", wintypes.DWORD), ("scheduling", wintypes.DWORD)]
    class IO(ctypes.Structure):
        _fields_ = [(name, ctypes.c_ulonglong) for name in ("read_ops", "write_ops", "other_ops", "read_bytes", "write_bytes", "other_bytes")]
    class ExtendedLimits(ctypes.Structure):
        _fields_ = [("basic", BasicLimits), ("io", IO), ("process_memory", ctypes.c_size_t), ("job_memory", ctypes.c_size_t),
                    ("peak_process", ctypes.c_size_t), ("peak_job", ctypes.c_size_t)]
    kernel = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel.CreateJobObjectW.restype = wintypes.HANDLE
    kernel.GetCurrentProcess.restype = wintypes.HANDLE
    kernel.SetInformationJobObject.argtypes = [wintypes.HANDLE, ctypes.c_int, ctypes.c_void_p, wintypes.DWORD]
    kernel.AssignProcessToJobObject.argtypes = [wintypes.HANDLE, wintypes.HANDLE]
    handle = kernel.CreateJobObjectW(None, None)
    limits = ExtendedLimits(); limits.basic.flags = 0x100; limits.process_memory = 512 * 1024 * 1024
    if not handle or not kernel.SetInformationJobObject(handle, 9, ctypes.byref(limits), ctypes.sizeof(limits)) or not kernel.AssignProcessToJobObject(handle, kernel.GetCurrentProcess()):
        raise ValueError("Parser memory limit unavailable")
    return handle


def parse_bytes(data: bytes, kind: str) -> str:
    if kind == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted or len(reader.pages) > 50:
            raise ValueError("Unsupported PDF")
        text = ""
        for page in reader.pages:
            content = page.get_contents()
            if content and len(content.get_data()) > 10 * 1024 * 1024:
                raise ValueError("PDF content too large")
            text += (page.extract_text() or "") + "\n"
            if len(text) > MAX_TEXT:
                break
        return text[:MAX_TEXT + 1]
    from docx import Document
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 300 or sum(e.file_size for e in entries) > 20 * 1024 * 1024:
            raise ValueError("Archive too large")
        for item in entries:
            if item.flag_bits & 1 or item.file_size / max(item.compress_size, 1) > 200 or "vbaproject" in item.filename.lower():
                raise ValueError("Unsupported archive")
            if item.filename.endswith((".xml", ".rels")):
                xml = archive.read(item)
                if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                    raise ValueError("External content is unsupported")
                if item.filename.endswith(".rels"):
                    from xml.etree import ElementTree
                    for relation in ElementTree.fromstring(xml):
                        if relation.get("TargetMode") == "External" and not relation.get("Type", "").endswith("/hyperlink"):
                            raise ValueError("External content is unsupported")
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)[:MAX_TEXT + 1]


if __name__ == "__main__":
    try:
        memory_limit_handle = bound_memory()
        print(json.dumps({"text": parse_bytes(sys.stdin.buffer.read(5 * 1024 * 1024 + 1), sys.argv[1])}))
    except Exception:
        sys.exit(2)
