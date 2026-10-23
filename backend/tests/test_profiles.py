"""Private journey regression checks use an isolated database and labelled demo documents."""
import io
import logging
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import create_engine, select, delete
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from backend.app.main import app
from backend.app.evidence_api import get_session
from backend.app.ingestion.runner import import_all
from backend.app.models import Base, PrivateSession, Profile, ProfileSkill, Role, Skill
from backend.app.skill_extraction import extract, normalize, parse_upload, provider_schema
from backend.app.profile_schemas import ExtractedSkills, SkillMention

HEADERS = {"X-Noka-CSRF": "1", "Origin": "http://testserver"}
MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture(scope="module")
def database():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    import_all(engine, Path(__file__).resolve().parents[2] / "data/fixtures", offline=True)
    yield engine
    engine.dispose()


@pytest.fixture
def client(database):
    def session():
        with Session(database) as db:
            yield db
    app.dependency_overrides[get_session] = session
    with Session(database) as db:
        for table in (ProfileSkill, Profile, PrivateSession):
            db.execute(delete(table))
        db.commit()
    with TestClient(app) as client:
        yield client
    app.dependency_overrides.clear()


def create(client):
    response = client.post('/api/profiles', json={"is_demo": True}, headers=HEADERS)
    assert response.status_code == 200
    return response.json()['id']


def payload(database):
    with Session(database) as db:
        skills = list(db.scalars(select(Skill).order_by(Skill.id)))[:5]
    return {"skills": [{"skill_id": s.id, "level": "strong", "years": None} for s in skills], "country": "MY", "interests": []}


def docx_bytes():
    doc = Document()
    doc.add_heading("Demo profile â€” illustrative experience", 0)
    doc.add_paragraph("Python programming, critical thinking, active listening, writing and mathematics.")
    out = io.BytesIO(); doc.save(out)
    return out.getvalue()


def pdf_bytes():
    writer = PdfWriter(); page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject(); stream.set_data(b'BT /F1 12 Tf 30 720 Td (Demo CV: Programming, Writing, Mathematics) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    out = io.BytesIO(); writer.write(out)
    return out.getvalue()


def test_cookie_is_opaque_and_ownership_protected(client, database):
    profile_id = create(client)
    cookie = client.cookies.get('noka_session')
    assert len(cookie) >= 40
    response = client.post('/api/profiles', json={}, headers=HEADERS)
    assert response.json()['id'] == profile_id
    with Session(database) as db:
        assert db.get(PrivateSession, cookie) is None
        assert db.get(Profile, profile_id).session_hash != cookie
    with TestClient(app) as other:
        for suffix in ('', '/matches'):
            assert other.get('/api/profiles/' + profile_id + suffix).status_code == 404
        assert other.put('/api/profiles/' + profile_id + '/skills', json=payload(database), headers=HEADERS).status_code == 404
        assert other.post('/api/profiles/' + profile_id + '/resume', data={'text': 'Programming'}, headers=HEADERS).status_code == 404
        assert other.delete('/api/profiles/' + profile_id, headers=HEADERS).status_code == 404
    fresh = TestClient(app)
    created = fresh.post('/api/profiles', json={}, headers=HEADERS)
    assert 'HttpOnly' in created.headers['set-cookie'] and 'SameSite=strict' in created.headers['set-cookie']
    assert 'no-store' in client.get('/api/profiles/current').headers['cache-control']


@pytest.mark.parametrize('headers', [{}, {'X-Noka-CSRF': '1', 'Origin': 'https://evil.example'}, {'X-Noka-CSRF': '1', 'Sec-Fetch-Site': 'cross-site'}])
def test_csrf_rejected(client, headers):
    assert client.post('/api/profiles', json={}, headers=headers).status_code == 403


def test_manual_only_deterministic_matches_and_deletion(client, database):
    profile_id = create(client)
    assert client.get(f'/api/profiles/{profile_id}/matches').json()['status'] == 'insufficient_profile_information'
    response = client.put(f'/api/profiles/{profile_id}/skills', json=payload(database), headers=HEADERS)
    assert response.status_code == 200
    first = client.get(f'/api/profiles/{profile_id}/matches').json()
    assert first == client.get(f'/api/profiles/{profile_id}/matches').json()
    # This fixture contains only general basic skills; do not force five recommendations.
    assert first['matches'] == []
    # An explicit target remains available even without a discovery connection.
    data=payload(database)
    data['interests']=['15-1252.00']
    assert client.put(f'/api/profiles/{profile_id}/skills', json=data, headers=HEADERS).status_code == 200
    first=client.get(f'/api/profiles/{profile_id}/matches').json()
    assert any(m['role_id']=='15-1252.00' and m['is_target'] for m in first['matches'])
    assert first == client.get(f'/api/profiles/{profile_id}/matches').json()
    assert all(0 <= m['score'] <= 100 and m['gaps'] and m['source_ids'] for m in first['matches'])
    assert all(m['is_target'] or m['discovery_relevance'] > 0 for m in first['matches'])
    assert 'not a probability' in first['methodology']
    assert client.delete(f'/api/profiles/{profile_id}', headers=HEADERS).status_code == 204
    assert client.get('/api/profiles/current').json() == {'profile': None}


def test_invalid_unknown_duplicate_and_partial_skills(client, database):
    profile_id = create(client); data = payload(database)
    data['skills'] = data['skills'][:1]
    assert client.put(f'/api/profiles/{profile_id}/skills', json=data, headers=HEADERS).status_code == 200
    assert client.get(f'/api/profiles/{profile_id}/matches').json()['status'] == 'insufficient_profile_information'
    data['skills'] *= 2
    assert client.put(f'/api/profiles/{profile_id}/skills', json=data, headers=HEADERS).status_code == 422
    data['skills'] = [{'skill_id': 'fictional', 'level': 'strong', 'years': None}]
    assert client.put(f'/api/profiles/{profile_id}/skills', json=data, headers=HEADERS).status_code == 422


@pytest.mark.parametrize('content,name,mime', [(b'evil', 'cv.exe', 'application/octet-stream'), (b'not pdf', 'cv.pdf', 'application/pdf'), (b'%PDF-bad', 'cv.pdf', MIME), (b'x' * (5*1024*1024+1), 'cv.pdf', 'application/pdf')], ids=['extension','magic','mime','oversize'])
def test_invalid_file(client, content, name, mime):
    profile_id = create(client)
    assert client.post(f'/api/profiles/{profile_id}/resume', files={'file': (name, content, mime)}, headers=HEADERS).status_code in (413, 422)


@pytest.mark.parametrize('factory,name,mime', [(docx_bytes,'demo.docx',MIME),(pdf_bytes,'demo.pdf','application/pdf')])
def test_real_document_parser_and_no_resume_storage(client, database, factory, name, mime):
    profile_id = create(client)
    response = client.post(f'/api/profiles/{profile_id}/resume', files={'file': (name, factory(), mime)}, headers=HEADERS)
    assert response.status_code == 200
    assert len(response.json()['candidates']) >= 3
    assert client.get(f'/api/profiles/{profile_id}').json()['skills'] == []
    assert not any('resume' in c.name or 'text' in c.name or 'name' in c.name for c in Profile.__table__.columns if c.name != 'analysis_context')
    with Session(database) as db:
        assert db.get(Profile, profile_id).analysis_context is None


def test_provider_failure_and_partial_output(client, database, monkeypatch, caplog):
    profile_id = create(client)
    def unavailable(text, skills):
        raise RuntimeError('Sensitive Name sensitive@example.com')
    monkeypatch.setattr('backend.app.skill_extraction.provider_extract', unavailable)
    with caplog.at_level(logging.WARNING):
        response = client.post(f'/api/profiles/{profile_id}/resume', data={'text': 'Sensitive Name sensitive@example.com Programming and Writing', 'use_ai': 'true'}, headers=HEADERS)
    assert response.json()['mode'] == 'manual_fallback'
    assert 'sensitive@example.com' not in response.text
    assert 'Sensitive Name' not in caplog.text and 'sensitive@example.com' not in caplog.text
    def invented(text, skills):
        return provider_schema(skills).model_validate({'skills':[{'mention':'Programming','years':None,'evidence_excerpt':'invented quote','confidence':'low','canonical_skill_ids':[s.id for s in skills if s.canonical_name=='Programming']}]})
    monkeypatch.setattr('backend.app.skill_extraction.provider_extract', invented)
    response = client.post(f'/api/profiles/{profile_id}/resume', data={'text': 'Programming', 'use_ai': 'true'}, headers=HEADERS)
    assert response.json()['mode'] == 'openai'
    assert response.json()['candidates'] == []
    assert response.json()['warnings']


def test_normalization_ambiguous_unknown_and_fuzzy(database):
    with Session(database) as db:
        skills = list(db.scalars(select(Skill)))
    assert normalize('Python',skills)[0] == 'alias'
    assert normalize('communication',skills)[0] == 'ambiguous'
    assert normalize('programmin',skills)[0] == 'fuzzy'
    assert normalize('teleportation',skills) == ('unmapped', [])
    assert not extract('No supported phrases here.',skills).candidates


def test_chunked_body_limit_and_strict_schema(client):
    profile_id = create(client)
    assert client.post('/api/profiles', content=iter([b'x'*1024*1024]*6), headers=HEADERS).status_code == 413
    assert client.post(f'/api/profiles/{profile_id}/resume', data={'text': 'x'*20001}, headers=HEADERS).status_code == 422
    assert client.post('/api/profiles', json={'name':'Never stored'}, headers=HEADERS).status_code == 422


def test_expired_session_is_not_authorized(client, database):
    profile_id = create(client)
    with Session(database) as db:
        owner = db.get(PrivateSession, db.get(Profile, profile_id).session_hash)
        owner.expires_at = datetime.now(timezone.utc) - timedelta(seconds=1)
        db.commit()
    assert client.get(f'/api/profiles/{profile_id}').status_code == 404
    assert client.get('/api/profiles/current').json()['profile'] is None


def test_score_bounds_interest_and_experience_independence(client, database):
    profile_id = create(client)
    with Session(database) as db:
        ids = list(db.scalars(select(Skill.id)))
        role = db.scalar(select(Role.id).order_by(Role.title))
    data = {'skills': [{'skill_id': i, 'level':'strong','years':0} for i in ids], 'country':'MY', 'interests':[]}
    client.put(f'/api/profiles/{profile_id}/skills', json=data, headers=HEADERS)
    full = client.get(f'/api/profiles/{profile_id}/matches').json()['matches']
    assert len(full) == 5 and all(m['score'] == 100 and not m['gaps'] for m in full)
    for skill in data['skills']: skill['level'] = 'working'
    client.put(f'/api/profiles/{profile_id}/skills', json=data, headers=HEADERS)
    base = client.get(f'/api/profiles/{profile_id}/matches').json()['matches']
    assert all(m['score'] == 65 for m in base)
    for skill in data['skills']: skill['years'] = 60
    data['country'] = 'SG'
    client.put(f'/api/profiles/{profile_id}/skills', json=data, headers=HEADERS)
    assert client.get(f'/api/profiles/{profile_id}/matches').json()['matches'] == base
    data['interests'] = [role]
    client.put(f'/api/profiles/{profile_id}/skills', json=data, headers=HEADERS)
    adjusted = client.get(f'/api/profiles/{profile_id}/matches').json()['matches']
    assert adjusted[0]['role_id'] == role and adjusted[0]['score'] == 65
    assert adjusted[0]['is_target'] and adjusted[0]['interest_adjustment'] == 0


@pytest.mark.parametrize('malicious', ['macro', 'external-template', 'entity', 'zip-bomb'])
def test_docx_active_content_rejected(malicious):
    source = zipfile.ZipFile(io.BytesIO(docx_bytes()))
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for item in source.infolist(): archive.writestr(item.filename, source.read(item))
        if malicious == 'macro': archive.writestr('word/vbaProject.bin', b'not executed')
        elif malicious == 'external-template': archive.writestr('word/_rels/template.rels', '<Relationships><Relationship TargetMode="External" Type="template" Target="https://evil.example"/></Relationships>')
        elif malicious == 'entity': archive.writestr('word/evil.xml', '<!DOCTYPE x [<!ENTITY e SYSTEM "file:///secret">]><x/>')
        else: archive.writestr('word/large.bin', b'x' * (21 * 1024 * 1024))
    with pytest.raises(ValueError): parse_upload(output.getvalue(), 'demo.docx', MIME)
