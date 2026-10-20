"""Create a labelled illustrative CV for the local parser demo; contains no real person."""
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]


def main():
    doc = Document()
    doc.core_properties.author = "Noka"
    doc.core_properties.title = "Demo profile — illustrative experience"
    doc.add_heading("Demo profile", 0)
    doc.add_paragraph("Illustrative experience for testing Noka. This is not a real person or a claim about a candidate.")
    doc.add_heading("Project experience", 1)
    doc.add_paragraph("Built a Python programming project. Practised critical thinking, complex problem solving, active learning and mathematics while investigating problems.")
    doc.add_paragraph("Used writing, speaking, active listening and coordination to explain ideas and organize a team project.")
    target = ROOT / "data/demo/demo_cv.docx"
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    print("Created labelled demo CV: data/demo/demo_cv.docx")


if __name__ == "__main__":
    main()
